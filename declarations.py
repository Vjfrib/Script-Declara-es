import os
from docx import Document

def declarations(dados):
    nome_disc = dados.get('nome_disc')
    cpf = dados.get('cpf')
    title = dados.get('title')
    area = dados.get('area')
    formato = dados.get('formato')

    # Caminho do arquivo de template
    caminho_template = os.path.join('template', 'Declarações $nomeDisc.docx')
    
    # Verifica se o arquivo de template existe
    if not os.path.exists(caminho_template):
        print(f"O arquivo {caminho_template} não foi encontrado.")
        return
    
    # Abre o arquivo de template
    doc = Document(caminho_template)
    
    # Substitui os placeholders pelos valores fornecidos sem alterar a formatação
    for paragrafo in doc.paragraphs:
        for run in paragrafo.runs:
            run.text = run.text.replace('$nomeDisc', nome_disc)
            run.text = run.text.replace('$cpf', cpf)
            run.text = run.text.replace('$title', title)
            run.text = run.text.replace('$area', area)
            run.text = run.text.replace('$formato', formato)
    
    # Cria a nova pasta se ela não existir
    nova_pasta = f'DEFESA {nome_disc}'
    if not os.path.exists(nova_pasta):
        os.makedirs(nova_pasta)
    
    # Define o caminho do novo arquivo
    novo_caminho_arquivo = os.path.join(nova_pasta, f'Declarações {nome_disc}.docx')
    
    # Salva o documento editado
    doc.save(novo_caminho_arquivo)
    
    print(f"Arquivo editado salvo como: {novo_caminho_arquivo}")
